#!/usr/bin/env python3
"""Gera o pitch deck do Lastro (6 slides, consolidado) com python-pptx.

Reescrito para trocar o pptxgenjs por python-pptx: a biblioteca anterior
emitia um chart XML com um eixo referenciado e nunca declarado (bug conhecido
em barras 2D), o que passava em todo validador disponível aqui mas o
PowerPoint de verdade recusava a abrir o arquivo inteiro. python-pptx não usa
esse gerador de chart (nem precisamos de chart nativo: a comparação de
orçamento é desenhada com formas simples), e o pacote OOXML que ele produz é
o caminho mais testado para compatibilidade real com o PowerPoint.
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.oxml.ns import qn

# ---------- palette ----------
WARM_BLACK = RGBColor(0x11, 0x10, 0x0F)
CARD_DARK = RGBColor(0x1B, 0x1A, 0x17)
SAND = RGBColor(0xF6, 0xF4, 0xEF)
GRAPHITE = RGBColor(0x5B, 0x64, 0x72)
VIOLETA = RGBColor(0xA1, 0x2A, 0xEB)
BROWN = RGBColor(0x50, 0x20, 0x03)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0x8A, 0x8A, 0x8A)
MUTED_LIGHT = RGBColor(0xC9, 0xC2, 0xB4)
CARD_BORDER = RGBColor(0xE6, 0xDF, 0xD2)
RATING_A = RGBColor(0x2E, 0x7D, 0x5B)
RATING_B = RGBColor(0x3F, 0x7C, 0xAC)
RATING_C = RGBColor(0xD0, 0x8C, 0x3A)
RATING_D = RGBColor(0xB2, 0x3A, 0x3A)

FONT_HEAD = "Cambria"
FONT_BODY = "Calibri"

SLIDE_W, SLIDE_H = 13.333, 7.5

prs = Presentation()
prs.slide_width = Inches(SLIDE_W)
prs.slide_height = Inches(SLIDE_H)
BLANK = prs.slide_layouts[6]


def new_slide(bg=SAND):
    slide = prs.slides.add_slide(BLANK)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = bg
    return slide


def _no_shadow(shape):
    shape.shadow.inherit = False


def rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, radius=None):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius is not None else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    _no_shadow(sh)
    if radius is not None:
        try:
            sh.adjustments[0] = radius
        except IndexError:
            pass
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    return sh


def oval(slide, x, y, w, h, fill=None, line=None, line_w=1.0):
    sh = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(w), Inches(h))
    _no_shadow(sh)
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
    return sh


def line(slide, x1, y1, x2, y2, color, width_pt=1.0, dash=None):
    conn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    _no_shadow(conn)
    conn.line.color.rgb = color
    conn.line.width = Pt(width_pt)
    if dash:
        d = conn.line._get_or_add_ln()
        pd = d.makeelement(qn('a:prstDash'), {'val': dash})
        d.append(pd)
    return conn


def _set_letter_spacing(run, pts):
    rPr = run._r.get_or_add_rPr()
    rPr.set('spc', str(int(pts * 100)))


def text(slide, x, y, w, h, s, size=12, color=WARM_BLACK, bold=False, italic=False,
         align=PP_ALIGN.LEFT, font=FONT_BODY, valign=MSO_ANCHOR.TOP, line_spacing=None,
         letter_spacing=None, wrap=True):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = valign
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    lines = s.split("\n")
    for i, ln in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if line_spacing:
            p.line_spacing = line_spacing
        r = p.add_run()
        r.text = ln
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.name = font
        r.font.color.rgb = color
        if letter_spacing:
            _set_letter_spacing(r, letter_spacing)
    return tb


def kicker(slide, s, x=0.7, y=0.5, w=8.0, color=VIOLETA):
    text(slide, x, y, w, 0.35, s.upper(), size=12.5, color=color, bold=True,
         font=FONT_BODY, letter_spacing=1.5)


def page_number(slide, n, dark):
    text(slide, SLIDE_W - 0.9, SLIDE_H - 0.5, 0.6, 0.3, f"{n:02d}", size=9,
         color=MUTED_LIGHT if dark else MUTED, align=PP_ALIGN.RIGHT)


def graph_motif(slide, x, y, scale=1.0, color=VIOLETA):
    pts = [(0, 0), (0.55 * scale, -0.35 * scale), (0.95 * scale, 0.15 * scale), (0.35 * scale, 0.45 * scale)]
    for i in range(len(pts) - 1):
        line(slide, x + pts[i][0], y + pts[i][1], x + pts[i + 1][0], y + pts[i + 1][1], color, 1.0)
    for i, (dx, dy) in enumerate(pts):
        r = 0.09 if i == 0 else 0.065
        oval(slide, x + dx - r / 2, y + dy - r / 2, r, r, fill=color)


# =========================================================
# Slide 1 - Capa
# =========================================================
s = new_slide(WARM_BLACK)
graph_motif(s, 10.7, 1.3, scale=2.2)
graph_motif(s, 1.0, 5.7, scale=1.4)
text(s, 0, 2.55, SLIDE_W, 1.2, "LASTRO", size=64, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, font=FONT_HEAD, letter_spacing=3)
text(s, SLIDE_W / 2 - 5, 3.8, 10, 0.5,
     "Inteligência de risco relacional e recuperação de capital no agro",
     size=16, color=MUTED_LIGHT, italic=True, align=PP_ALIGN.CENTER)
rect(s, SLIDE_W / 2 - 1.1, 4.5, 2.2, 0.02, fill=VIOLETA)
text(s, SLIDE_W / 2 - 5, 4.7, 10, 0.4, "Hackathon PMI-DF 2026  ·  Case Krilltech",
     size=13, color=WHITE, align=PP_ALIGN.CENTER)
text(s, SLIDE_W / 2 - 5, 5.1, 10, 0.4, "Apresentação: Erick Mendes  ·  Equipe Lastro",
     size=11.5, color=MUTED_LIGHT, align=PP_ALIGN.CENTER)

# =========================================================
# Slide 2 - A tese + o diagnostico (combinado)
# =========================================================
s = new_slide(WARM_BLACK)
kicker(s, "A tese", y=0.45)
text(s, 0.9, 0.95, 11.4, 0.55,
     "Depois do pedido de Recuperação Judicial, não existe boa decisão.",
     size=25, color=WHITE, bold=True, font=FONT_HEAD, line_spacing=1.1)
text(s, 0.9, 1.65, 11.4, 0.5, "Só existe decisão antes.",
     size=25, color=VIOLETA, bold=True, font=FONT_HEAD)
text(s, 0.9, 2.35, 11, 0.35,
     "180 dias de stay period (Lei 11.101/2005, estendida ao produtor rural pela Lei 14.112/2020)",
     size=10.5, color=MUTED_LIGHT, italic=True)

kicker(s, "O diagnóstico", y=2.95, color=VIOLETA)
stats = [
    ("1.990", "pedidos de RJ no agro em 2025", "+56,4% sobre 2024", "Serasa Experian"),
    ("2,7% → 7,3%", "inadimplência do produtor rural PF", "em 12 meses", "Banco Central"),
    ("8,8%", "inadimplência rural no 1T2026", "recorde da série", "Serasa Experian"),
]
card_w, gap = 3.55, 0.35
start_x = (SLIDE_W - (card_w * 3 + gap * 2)) / 2
card_y, card_h = 3.4, 2.35
for i, (big, sub, tag, src) in enumerate(stats):
    x = start_x + i * (card_w + gap)
    rect(s, x, card_y, card_w, card_h, fill=CARD_DARK, line=RGBColor(0x2E, 0x2C, 0x27), line_w=0.75, radius=0.06)
    text(s, x + 0.2, card_y + 0.2, card_w - 0.4, 0.55, big, size=22, color=VIOLETA, bold=True,
         font=FONT_HEAD, align=PP_ALIGN.CENTER)
    text(s, x + 0.2, card_y + 0.85, card_w - 0.4, 0.6, sub, size=11, color=WHITE, align=PP_ALIGN.CENTER,
         line_spacing=1.1)
    text(s, x + 0.2, card_y + 1.5, card_w - 0.4, 0.35, tag, size=10.5, color=RGBColor(0xC9, 0x9C, 0x5A),
         bold=True, align=PP_ALIGN.CENTER)
    text(s, x + 0.2, card_y + 1.95, card_w - 0.4, 0.3, src, size=9, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
text(s, 0.9, 6.0, 11.4, 0.4,
     "A Krilltech vende a prazo, direto ao produtor, em 18 estados. A carteira dela está exatamente nessa população.",
     size=11.5, color=MUTED_LIGHT, italic=True, align=PP_ALIGN.CENTER)
page_number(s, 2, True)

# =========================================================
# Slide 3 - Como o Lastro funciona (motor + agentes + governanca)
# =========================================================
s = new_slide(SAND)
kicker(s, "Como o Lastro funciona", color=BROWN)
text(s, 0.7, 0.95, 12, 0.55, "Rede de risco, dois canais, quatro agentes",
     size=22, color=WARM_BLACK, bold=True, font=FONT_HEAD)

colY, colH = 1.75, 4.15
leftX, leftW = 0.7, 6.0
rightX, rightW = 7.0, 5.65

rect(s, leftX, colY, leftW, colH, fill=WHITE, line=CARD_BORDER, line_w=0.75, radius=0.045)
text(s, leftX + 0.3, colY + 0.25, leftW - 0.6, 0.35, "Motor de exposição — dois canais",
     size=14, color=WARM_BLACK, bold=True)
text(s, leftX + 0.3, colY + 0.75, leftW - 0.6, 0.3, "ESTRUTURAL · contamina por vínculo jurídico/patrimonial",
     size=10.5, color=BROWN, bold=True)
text(s, leftX + 0.3, colY + 1.08, leftW - 0.6, 0.45,
     "Grupo econômico 0,90  ·  Avalista em comum 0,85  ·  Sócio em comum (QSA) 0,70",
     size=11.5, color=WARM_BLACK, line_spacing=1.15)
text(s, leftX + 0.3, colY + 1.65, leftW - 0.6, 0.3, "SISTÊMICO · ninguém contamina; todos sofrem a mesma causa",
     size=10.5, color=VIOLETA, bold=True)
text(s, leftX + 0.3, colY + 1.98, leftW - 0.6, 0.45,
     "Região + cultura 0,75*  ·  Mesma revenda 0,65  ·  Mesma cultura 0,40",
     size=11.5, color=WARM_BLACK, line_spacing=1.15)
text(s, leftX + 0.3, colY + 2.55, leftW - 0.6, 0.55,
     "* peso cheio só com evento regional confirmando o choque (quebra de safra, alerta ZARC, queda de preço em até 365 dias); sem evento, cai para 0,30.",
     size=9.5, color=MUTED, italic=True, line_spacing=1.15)
text(s, leftX + 0.3, colY + 3.15, leftW - 0.6, 0.4, "Toda exposição guarda o caminho que a gerou.",
     size=11, color=WARM_BLACK, bold=True)

rect(s, rightX, colY, rightW, colH, fill=WHITE, line=CARD_BORDER, line_w=0.75, radius=0.045)
text(s, rightX + 0.3, colY + 0.25, rightW - 0.6, 0.35, "Onde entra IA de verdade — 4 agentes",
     size=14, color=WARM_BLACK, bold=True)
agents = [
    ("1", "Coletor & Parser", "lê base pública e PDF, devolve evento com fonte"),
    ("2", "Risco Agro & Climático", "cruza CAR, ZARC, cultura e safra com a região"),
    ("3", "Decisão & Scoring", "dispara o cálculo e interpreta — não calcula"),
    ("4", "Sintetizador", "redige o parecer citando só números do dossiê"),
]
ay = colY + 0.85
for num, title, desc in agents:
    oval(s, rightX + 0.3, ay, 0.32, 0.32, fill=VIOLETA)
    text(s, rightX + 0.3, ay, 0.32, 0.32, num, size=12, color=WHITE, bold=True,
         align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    text(s, rightX + 0.75, ay - 0.03, rightW - 1.05, 0.32, title, size=12, color=WARM_BLACK, bold=True)
    text(s, rightX + 0.75, ay + 0.28, rightW - 1.05, 0.32, desc, size=10, color=GRAPHITE, line_spacing=1.1)
    ay += 0.62
text(s, rightX + 0.3, colY + colH - 0.45, rightW - 0.6, 0.4,
     "IA entra onde há texto e ambiguidade. Onde há conta, é query determinística.",
     size=10, color=BROWN, italic=True)

rect(s, 0.7, colY + colH + 0.25, leftW + rightW + 0.3, 0.65, fill=WARM_BLACK, radius=0.06)
text(s, 1.0, colY + colH + 0.25, leftW + rightW - 0.3, 0.65,
     "O Lastro recomenda, nunca decide sozinho  ·  nunca prescreve instrumento jurídico  ·  nunca promete prever falência — mede exposição compartilhada",
     size=10.5, color=WHITE, bold=True, valign=MSO_ANCHOR.MIDDLE)
page_number(s, 3, False)

# =========================================================
# Slide 4 - A demonstracao + a decisao
# =========================================================
s = new_slide(WARM_BLACK)
kicker(s, "A demonstração")
text(s, 0.7, 0.85, 8.6, 0.5, "Um clique. Quatro clientes acendem. Um continua verde.",
     size=19, color=WHITE, bold=True, font=FONT_HEAD)
text(s, 0.7, 1.35, 3.6, 0.6, "R$ 1.258.000", size=24, color=VIOLETA, bold=True, font=FONT_HEAD)
text(s, 4.15, 1.4, 4.7, 0.55, "de exposição identificada\nantes do vencimento",
     size=10.5, color=MUTED_LIGHT, line_spacing=1.1)

gx0, gy0, gx1, gy1 = 0.6, 2.1, 8.9, 7.1
center = (2.5, 4.85)
nodes = [
    ("Terra Nova", "grupo econômico · A", "R$ 410 mil", (6.3, 2.65), 0.5, RATING_A, True),
    ("Fazenda Santa Luzia", "sócio comum + região · B", "R$ 520 mil", (8.15, 3.55), 0.42, RATING_B, True),
    ("Agropecuária Horizonte", "avalista comum · C", "R$ 230 mil", (8.0, 5.35), 0.38, RATING_C, False),
    ("Sítio Boa Esperança", "região/cultura (Conab) · C", "R$ 98 mil", (6.15, 6.35), 0.34, RATING_C, False),
]
for _, _, _, (nx, ny), _, _, _ in nodes:
    line(s, center[0], center[1], nx, ny, VIOLETA, 1.1, dash="sysDot")

oval(s, center[0] - 0.62, center[1] - 0.62, 1.24, 1.24, fill=RATING_D, line=WHITE, line_w=1.75)
text(s, center[0] - 0.62, center[1] - 0.33, 1.24, 0.55, "Agro Vale\ndo Cerrado", size=10.5, color=WHITE,
     bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, line_spacing=1.0)
text(s, center[0] - 1.1, center[1] + 0.68, 2.2, 0.3, "pediu RJ · rating D", size=9.5, color=MUTED_LIGHT,
     align=PP_ALIGN.CENTER)

for label, sub, val, (nx, ny), r, color, above in nodes:
    oval(s, nx - r, ny - r, r * 2, r * 2, fill=color, line=WHITE, line_w=1.25)
    label_h = 0.75
    label_y = ny - r - label_h - 0.05 if above else ny + r + 0.06
    text(s, nx - 1.25, label_y, 2.5, label_h, f"{label}\n{sub}\n{val}", size=9, color=WHITE,
         align=PP_ALIGN.CENTER, line_spacing=1.05)

oval(s, 1.15 - 0.32, 6.15 - 0.32, 0.64, 0.64, fill=RATING_A, line=WHITE, line_w=1.25)
text(s, 1.15 - 1.35, 6.15 + 0.4, 2.7, 0.55, "Ipê Amarelo\ncontrole · A · sem vínculo", size=9,
     color=MUTED_LIGHT, align=PP_ALIGN.CENTER, line_spacing=1.05)

panel_x, panel_y, panel_w, panel_h = 9.25, 2.1, 3.45, 4.55
rect(s, panel_x, panel_y, panel_w, panel_h, fill=CARD_DARK, line=RGBColor(0x2E, 0x2C, 0x27), line_w=0.75, radius=0.05)
text(s, panel_x + 0.3, panel_y + 0.3, panel_w - 0.6, 0.35, "A decisão", size=11.5, color=VIOLETA, bold=True,
     letter_spacing=1)
text(s, panel_x + 0.3, panel_y + 0.75, panel_w - 0.6, 0.4, "Sítio Boa Esperança", size=14, color=WHITE, bold=True)
text(s, panel_x + 0.3, panel_y + 1.2, panel_w - 0.6, 0.35, "Score 530 · Rating C · adimplente hoje",
     size=11, color=MUTED_LIGHT)
text(s, panel_x + 0.3, panel_y + 1.75, panel_w - 0.6, 0.9,
     "Vínculo: região e cultura com quebra de safra confirmada pela Conab",
     size=10.5, color=WHITE, line_spacing=1.2)
text(s, panel_x + 0.3, panel_y + 2.7, panel_w - 0.6, 0.9,
     "Recomendação: reduzir limite e condicionar a próxima venda, com prazo definido",
     size=10.5, color=WHITE, line_spacing=1.2)
text(s, panel_x + 0.3, panel_y + panel_h - 0.55, panel_w - 0.6, 0.4,
     "Decisão final: comitê de crédito", size=9.5, color=MUTED_LIGHT, italic=True)
page_number(s, 4, True)

# =========================================================
# Slide 5 - Plano de conclusao (equipe + cronograma + orcamento)
# =========================================================
s = new_slide(SAND)
kicker(s, "Plano de conclusão", color=BROWN)
text(s, 0.7, 0.95, 12, 0.55, "Da demonstração à produção: equipe, prazo e investimento",
     size=20, color=WARM_BLACK, bold=True, font=FONT_HEAD)
text(s, 0.7, 1.5, 12, 0.35, "Horizonte de 6 meses, do fim do hackathon até uso ativo pelo comitê de crédito",
     size=11, color=GRAPHITE, italic=True)

colY, colH = 2.05, 4.5
col_w, col_gap = 3.85, 0.3
col_x0 = (SLIDE_W - (col_w * 3 + col_gap * 2)) / 2

# Coluna 1: equipe
x = col_x0
rect(s, x, colY, col_w, colH, fill=WHITE, line=CARD_BORDER, line_w=0.75, radius=0.05)
text(s, x + 0.28, colY + 0.28, col_w - 0.56, 0.35, "Equipe", size=14, color=WARM_BLACK, bold=True)
team_lines = [
    "Tech Lead / Arquiteto (integral, 6 meses)",
    "2 Engenheiros de Dados / Backend (integral)",
    "Engenheiro Frontend (integral, 4 meses)",
    "Analista de Fontes Públicas (integral → parcial)",
    "Especialista de Domínio (consultoria pontual)",
    "Ponto focal da Krilltech (meio período, mês 2+)",
]
ty = colY + 0.8
for t in team_lines:
    oval(s, x + 0.3, ty + 0.07, 0.06, 0.06, fill=VIOLETA)
    text(s, x + 0.48, ty, col_w - 0.76, 0.5, t, size=10.5, color=WARM_BLACK, line_spacing=1.1)
    ty += 0.53

# Coluna 2: cronograma
x = col_x0 + (col_w + col_gap)
rect(s, x, colY, col_w, colH, fill=WHITE, line=CARD_BORDER, line_w=0.75, radius=0.05)
text(s, x + 0.28, colY + 0.28, col_w - 0.56, 0.35, "Cronograma", size=14, color=WARM_BLACK, bold=True)
phases = [
    ("Mês 1", "Conexão e fundação", "base real da Krilltech, mapeamento de vínculos via QSA"),
    ("Meses 2–3", "Automação e modo sombra", "coleta contínua; sistema roda ao lado do processo atual"),
    ("Meses 4–5", "Validação", "comparação dos alertas; calibração em 6, 12 e 24 meses"),
    ("Mês 6", "Virada para produção", "uso ativo pelo comitê de crédito; treinamento das equipes"),
]
py = colY + 0.8
for m, t, d in phases:
    text(s, x + 0.28, py, col_w - 0.56, 0.24, m.upper(), size=9.5, color=VIOLETA, bold=True, letter_spacing=0.5)
    text(s, x + 0.28, py + 0.24, col_w - 0.56, 0.28, t, size=11.5, color=WARM_BLACK, bold=True)
    text(s, x + 0.28, py + 0.53, col_w - 0.56, 0.4, d, size=9, color=GRAPHITE, line_spacing=1.1)
    py += 0.85

# Coluna 3: investimento
x = col_x0 + 2 * (col_w + col_gap)
rect(s, x, colY, col_w, colH, fill=WARM_BLACK, radius=0.05)
text(s, x + 0.28, colY + 0.28, col_w - 0.56, 0.35, "Investimento", size=14, color=WHITE, bold=True)
text(s, x + 0.28, colY + 0.8, col_w - 0.56, 0.6, "≈ R$ 366 mil", size=24, color=VIOLETA, bold=True, font=FONT_HEAD)
text(s, x + 0.28, colY + 1.45, col_w - 0.56, 0.75,
     "para os 6 meses até produção (equipe + infraestrutura)", size=10.5, color=WHITE, line_spacing=1.15)
text(s, x + 0.28, colY + 2.15, col_w - 0.56, 0.3, "Equipe R$ 331 mil  ·  Infra R$ 35 mil",
     size=10, color=MUTED_LIGHT)
line(s, x + 0.28, colY + 2.6, x + col_w - 0.28, colY + 2.6, RGBColor(0x3A, 0x38, 0x33), 0.75)
text(s, x + 0.28, colY + 2.75, col_w - 0.56, 0.45, "R$ 1.258.000", size=16, color=WHITE, bold=True, font=FONT_HEAD)
text(s, x + 0.28, colY + 3.2, col_w - 0.56, 0.75,
     "em exposição já identificada só no cenário de demonstração — já cobre o investimento do semestre",
     size=9.5, color=MUTED_LIGHT, line_spacing=1.1)

text(s, 0.7, colY + colH + 0.2, 12, 0.35, "Ordem de grandeza, hipótese declarada — não cotação fechada.",
     size=9.5, color=MUTED, italic=True, align=PP_ALIGN.CENTER)
page_number(s, 5, False)

# =========================================================
# Slide 6 - Fecho
# =========================================================
s = new_slide(WARM_BLACK)
graph_motif(s, 1.0, 1.2, scale=1.9)
graph_motif(s, 10.9, 5.6, scale=1.5)
text(s, 0, 2.5, SLIDE_W, 0.7, "O mesmo negócio.", size=32, color=WHITE, bold=True,
     align=PP_ALIGN.CENTER, font=FONT_HEAD)
text(s, 0, 3.25, SLIDE_W, 0.7, "Fechando o ciclo.", size=32, color=VIOLETA, bold=True,
     align=PP_ALIGN.CENTER, font=FONT_HEAD)
text(s, SLIDE_W / 2 - 5, 4.25, 10, 0.7,
     "A Krilltech otimiza a lavoura do cliente. O Lastro entrega a mesma inteligência sobre a saúde financeira de quem compra dela.",
     size=13, color=MUTED_LIGHT, italic=True, align=PP_ALIGN.CENTER, line_spacing=1.2)
rect(s, SLIDE_W / 2 - 1.1, 5.35, 2.2, 0.02, fill=VIOLETA)
text(s, 0, 5.55, SLIDE_W, 0.5, "LASTRO", size=20, color=WHITE, bold=True, align=PP_ALIGN.CENTER,
     font=FONT_HEAD, letter_spacing=2.5)
text(s, 0, 6.2, SLIDE_W, 0.4, "Obrigado · Equipe Lastro · Hackathon PMI-DF 2026",
     size=11, color=MUTED, align=PP_ALIGN.CENTER)

# ---------- speaker notes ----------
notes = [
    "Abertura. Deixar o wordmark no ar por 1-2 segundos antes de começar a falar.",
    "Tese: quando um cliente pede RJ começa o stay period de 180 dias. Pausa de 2s depois de 'Só existe decisão antes'. Diagnóstico: 1.990 RJs (+56,4%), inadimplência de 2,7% a 7,3%, Krilltech vende em 18 estados.",
    "Como funciona (apoio, não precisa ler tudo em voz alta): motor de dois canais, os quatro agentes, e a regra de que o sistema recomenda mas nunca decide.",
    "A demonstração: um clique na Agro Vale do Cerrado propaga o risco e quatro clientes acendem, nenhum em atraso hoje. O Ipê Amarelo continua verde: sem vínculo, sem contaminação. A decisão: Sítio Boa Esperança, score 530, rating C, recomendação com prazo. Quem decide é o comitê.",
    "Fecho concreto: equipe, cronograma de 6 meses e investimento de ordem de R$ 366 mil, contra R$ 1,258 milhão já identificados na própria demo.",
    "Fecho. Não cortar esta fala mesmo se o tempo apertar.",
]
for slide, note in zip(prs.slides, notes):
    slide.notes_slide.notes_text_frame.text = note

OUT = "/mnt/user-data/outputs/Lastro-Pitch.pptx"
prs.save(OUT)
print("OK:", OUT)
